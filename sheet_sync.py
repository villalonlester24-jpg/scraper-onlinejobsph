#!/usr/bin/env python3
"""
Push the newest rows of output/jobs.xlsx to a Google Sheets web app.

    python sheet_sync.py --rows 60

Reads the workbook the scraper just wrote, takes the last N rows, and POSTs
them to the Apps Script endpoint in apps-script/Code.gs. The script on the
other end skips any job_id it already has, so:

  * sending the same row twice is harmless, and
  * every run repairs the last N rows, so a Sheets outage costs one run of
    delay rather than a permanent hole in the sheet.

That second property is the whole reason this exists as a separate step
instead of using the scraper's --webhook. --webhook can only send jobs that
are new *to the workbook*, so anything it fails to deliver is already recorded
locally and never offered again.

Config comes from the environment, matching the GitHub Action:

    SHEETS_WEBHOOK_URL     the /exec URL from the Apps Script deployment
    SHEETS_WEBHOOK_TOKEN   the token set in Script properties

Both can be overridden with --url / --token.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
from pathlib import Path

import requests
from openpyxl import load_workbook

# The Apps Script builds its own columns, so these are just the workbook header
# names to pull from — kept as constants so a rename fails loudly here rather
# than silently sending blanks.
H = {
    "job_id": "Job ID",
    "title": "Title",
    "wage": "Wage / Salary",
    "job_type": "Job Type",
    "posted": "Posted",
    "found_at": "Scraped At",
    "tags": "Tags",
    "url": "URL",
    "description": "Job Overview (full)",
}


def read_recent(path: Path, rows: int) -> tuple[list[dict], int]:
    """Return (records, total_rows) for the last `rows` rows of the workbook."""
    if not path.exists():
        raise SystemExit(f"no workbook at {path} — run the scraper first")

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    table = list(ws.iter_rows(values_only=True))
    wb.close()

    if not table:
        return [], 0

    header = ["" if h is None else str(h) for h in table[0]]
    missing = [name for name in H.values() if name not in header]
    if missing:
        raise SystemExit(f"{path} is missing column(s): {', '.join(missing)}")

    index = {key: header.index(name) for key, name in H.items()}

    data = [r for r in table[1:] if any(v is not None for v in r)]
    total = len(data)
    window = data[-rows:] if rows > 0 else data

    records = []
    for row in window:
        rec = {}
        for key, col in index.items():
            value = row[col] if col < len(row) else None
            rec[key] = "" if value is None else str(value)
        if rec["job_id"]:
            records.append(rec)
    return records, total


def post(url: str, token: str | None, payload: dict, timeout: int = 180) -> dict:
    """POST to the Apps Script web app and return its JSON reply.

    Apps Script answers a POST with a 302 to a one-time googleusercontent URL.
    requests follows it and downgrades to GET, which is correct here — the
    script has already run and the redirect body carries its result.
    """
    endpoint = url
    if token:
        endpoint = f"{url}{'&' if '?' in url else '?'}token={token}"

    resp = requests.post(
        endpoint,
        json=payload,
        headers={"Content-Type": "application/json"},
        timeout=timeout,
    )
    if resp.status_code >= 300:
        raise SystemExit(
            f"web app returned HTTP {resp.status_code}: {resp.text[:400]}"
        )
    try:
        return resp.json()
    except ValueError:
        raise SystemExit(
            "web app did not return JSON — is the deployment set to "
            f"'Anyone'? First 400 chars:\n{resp.text[:400]}"
        )


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    p.add_argument("--workbook", type=Path, default=Path("output/jobs.xlsx"),
                   help="workbook to read (default output/jobs.xlsx)")
    p.add_argument("--rows", type=int, default=60,
                   help="how many of the newest rows to send (default 60, 0 = all)")
    p.add_argument("--url", default=os.environ.get("SHEETS_WEBHOOK_URL"),
                   help="Apps Script /exec URL (default $SHEETS_WEBHOOK_URL)")
    p.add_argument("--token", default=os.environ.get("SHEETS_WEBHOOK_TOKEN"),
                   help="shared token (default $SHEETS_WEBHOOK_TOKEN)")
    p.add_argument("--dry-run", action="store_true",
                   help="read and report, but don't send anything")
    args = p.parse_args()

    records, total = read_recent(args.workbook, args.rows)
    print(f"{args.workbook}: {total} rows, sending the newest {len(records)}")

    if not records:
        print("nothing to send")
        return 0

    if args.dry_run:
        print(json.dumps(records[0], indent=2)[:800])
        return 0

    if not args.url:
        print("no web app URL — set SHEETS_WEBHOOK_URL or pass --url",
              file=sys.stderr)
        return 1

    reply = post(args.url, args.token, {
        "run_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "rows": records,
    })

    if not reply.get("ok"):
        print(f"  ! web app refused it: {reply.get('error', reply)}",
              file=sys.stderr)
        return 1

    print(f"  added {reply.get('added', 0)}, "
          f"already there {reply.get('skipped', 0)}, "
          f"sheet now holds {reply.get('total', '?')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
